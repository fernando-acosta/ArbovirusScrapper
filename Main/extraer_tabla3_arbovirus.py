#!/usr/bin/env python3
"""Extrae Tabla 3 de informes PDF de arbovirus y crea hojas de cálculo."""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections.abc import Iterable
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    sys.exit(
        "Faltan dependencias. Instálalas con: "
        "python -m pip install -r requirements-extract.txt"
    )


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_PDF_DIR = SCRIPT_DIR / "informes_arbovirus_PR" / "PDFs"
DEFAULT_INVENTORY = SCRIPT_DIR / "informes_arbovirus_PR" / "inventario.csv"
DEFAULT_OUTPUT_DIR = SCRIPT_DIR / "tabla3_arbovirus_PR"
HEADERS = [
    "Fecha de emisión",
    "Fuente (URL)",
    "Municipio",
    "Frecuencia",
    "Porciento",
]
MONTHS = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "setiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}
TABLE_ROW_RE = re.compile(
    r"^\s*(?P<municipio>.+?)\s+"
    r"(?P<frecuencia>\d[\d,]*(?:\s*\([+-]?\d+\))?[*†‡]?)\s+"
    r"(?P<porciento>\d+(?:[.,]\d+)?)\s*%?\s*$"
)
DATE_SPANISH_RE = re.compile(
    r"(?<!\w)(?P<day>\d{1,2})\s+de\s+(?P<month>[a-z]+)\s+de\s+"
    r"(?P<year>\d{4})(?!\w)"
)
DATE_NUMERIC_RE = re.compile(
    r"(?<!\d)(?P<day>\d{1,2})[/-](?P<month>\d{1,2})[/-](?P<year>\d{4})(?!\d)"
)
DATE_ISO_RE = re.compile(r"(?<!\d)(?P<year>\d{4})-(?P<month>\d{1,2})-(?P<day>\d{1,2})(?!\d)")
TABLE3_RE = re.compile(r"\btabla\s*3\b", re.IGNORECASE)
NEXT_TABLE_RE = re.compile(r"^\s*tabla\s*[4-9]\b", re.IGNORECASE)


def normalize(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text)
    return "".join(char for char in normalized if not unicodedata.combining(char)).lower()


def parse_date_match(match: re.Match[str], spanish_month: bool = False) -> date | None:
    try:
        day = int(match.group("day"))
        month_text = match.group("month")
        month = MONTHS.get(normalize(month_text)) if spanish_month else int(month_text)
        year = int(match.group("year"))
        if month is None:
            return None
        return date(year, month, day)
    except (ValueError, TypeError):
        return None


def parse_issue_date(text: str) -> date | None:
    normalized = normalize(text)
    labeled_date = re.search(
        r"\bfecha(?:\s+de\s+(?:emision|publicacion))?\s*[:\-]?\s*(.{0,100})",
        normalized,
    )
    candidates = [labeled_date.group(1)] if labeled_date else []
    candidates.append(normalized)
    for candidate in candidates:
        for match in DATE_SPANISH_RE.finditer(candidate):
            result = parse_date_match(match, spanish_month=True)
            if result:
                return result
        for pattern in (DATE_NUMERIC_RE, DATE_ISO_RE):
            for match in pattern.finditer(candidate):
                result = parse_date_match(match)
                if result:
                    return result
    return None


def parse_table3(pages: Iterable[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    started = False
    header_seen = False
    for page_text in pages:
        for line in page_text.splitlines():
            stripped = line.strip()
            if TABLE3_RE.search(stripped):
                started = True
            if not started:
                continue
            if NEXT_TABLE_RE.match(stripped):
                return rows
            normalized_line = normalize(stripped)
            if all(term in normalized_line for term in ("municipio", "frecuencia", "porciento")):
                header_seen = True
                continue
            if not header_seen:
                continue
            match = TABLE_ROW_RE.match(stripped)
            if not match:
                continue
            municipality = match.group("municipio").strip()
            if normalize(municipality) == "total":
                return rows
            frequency_text = match.group("frecuencia").replace(",", "")
            frequency_text = re.sub(r"\s+", " ", frequency_text).strip()
            frequency: int | str
            if re.fullmatch(r"\d+", frequency_text):
                frequency = int(frequency_text)
            else:
                frequency = frequency_text
            percent = float(match.group("porciento").replace(",", "."))
            rows.append(
                {
                    "Municipio": municipality,
                    "Frecuencia": frequency,
                    "Porciento": percent,
                }
            )
    return rows


def get_issue_date(pdf: pdfplumber.PDF) -> date | None:
    cover_text = "\n".join((page.extract_text() or "") for page in pdf.pages[:3])
    return parse_issue_date(cover_text)


def read_source_urls(inventory_path: Path) -> dict[str, str]:
    if not inventory_path.is_file():
        raise FileNotFoundError(f"No se encontró el inventario: {inventory_path}")
    sources: dict[str, str] = {}
    with inventory_path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        for row in csv.DictReader(csv_file):
            filename = (row.get("archivo") or "").strip()
            url = (row.get("url") or "").strip()
            if filename and url and filename not in sources:
                sources[filename] = url
    return sources


def add_worksheet(workbook: Workbook, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.active
    sheet.title = "Tabla 3"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sheet.append([row[header] for header in HEADERS])
        date_cell = sheet.cell(sheet.max_row, 1)
        date_cell.number_format = "dd/mm/yyyy"
        source_cell = sheet.cell(sheet.max_row, 2)
        source_cell.hyperlink = source_cell.value
        source_cell.style = "Hyperlink"
        sheet.cell(sheet.max_row, 5).number_format = "0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [18, 62, 25, 18, 15]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_outputs(
    output_dir: Path,
    rows: list[dict[str, Any]],
    errors: list[dict[str, str]],
    summary: dict[str, Any],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = output_dir / "tabla3_arbovirus_PR.xlsx"
    workbook = Workbook()
    add_worksheet(workbook, rows)
    workbook.save(xlsx_path)

    csv_path = output_dir / "tabla3_arbovirus_PR.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=HEADERS)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    **row,
                    "Fecha de emisión": row["Fecha de emisión"].strftime("%d/%m/%Y"),
                }
            )

    errors_path = output_dir / "errores_extraccion.csv"
    with errors_path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.DictWriter(
            csv_file, fieldnames=["archivo", "url", "problema"]
        )
        writer.writeheader()
        writer.writerows(errors)

    (output_dir / "resumen.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extraer la Tabla 3 de los informes PDF de arbovirus."
    )
    parser.add_argument("--pdf-dir", type=Path, default=DEFAULT_PDF_DIR)
    parser.add_argument("--inventario", type=Path, default=DEFAULT_INVENTORY)
    parser.add_argument("--salida", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--inicio", type=date.fromisoformat, default=date(2023, 2, 1))
    parser.add_argument("--fin", type=date.fromisoformat, default=date(2026, 10, 31))
    args = parser.parse_args()
    if args.inicio > args.fin:
        parser.error("--inicio no puede ser posterior a --fin")
    if not args.pdf_dir.is_dir():
        parser.error(f"No se encontró la carpeta de PDFs: {args.pdf_dir}")

    source_urls = read_source_urls(args.inventario)
    pdf_paths = sorted(args.pdf_dir.glob("*.pdf"))
    rows: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    outside_range = 0
    missing_dates = 0
    missing_tables = 0

    for index, pdf_path in enumerate(pdf_paths, start=1):
        source_url = source_urls.get(pdf_path.name, "")
        try:
            with pdfplumber.open(pdf_path) as pdf:
                issue_date = get_issue_date(pdf)
                if issue_date is None:
                    missing_dates += 1
                    errors.append(
                        {
                            "archivo": pdf_path.name,
                            "url": source_url,
                            "problema": "No se pudo identificar la fecha de emisión en las primeras páginas.",
                        }
                    )
                    continue
                if not args.inicio <= issue_date <= args.fin:
                    outside_range += 1
                    continue
                table_rows = parse_table3(
                    (page.extract_text() or "") for page in pdf.pages
                )
        except Exception as exc:
            errors.append(
                {
                    "archivo": pdf_path.name,
                    "url": source_url,
                    "problema": f"Error al leer el PDF: {exc}",
                }
            )
            continue

        if not table_rows:
            missing_tables += 1
            errors.append(
                {
                    "archivo": pdf_path.name,
                    "url": source_url,
                    "problema": "No se encontró la Tabla 3 con filas legibles.",
                }
            )
            continue
        if not source_url:
            errors.append(
                {
                    "archivo": pdf_path.name,
                    "url": "",
                    "problema": "No se encontró un enlace de origen para este PDF en el inventario.",
                }
            )
        rows.extend(
            {
                "Fecha de emisión": issue_date,
                "Fuente (URL)": source_url,
                **table_row,
            }
            for table_row in table_rows
        )
        if index % 25 == 0:
            print(f"Procesados {index}/{len(pdf_paths)} PDF...", flush=True)

    rows.sort(key=lambda row: (row["Fecha de emisión"], row["Municipio"]))
    summary = {
        "fecha_generacion": datetime.now().isoformat(timespec="seconds"),
        "rango_inclusivo": {
            "inicio": args.inicio.isoformat(),
            "fin": args.fin.isoformat(),
        },
        "pdfs_en_carpeta": len(pdf_paths),
        "pdfs_fuera_del_rango": outside_range,
        "pdfs_sin_fecha_identificable": missing_dates,
        "pdfs_sin_tabla3_legible": missing_tables,
        "filas_extraidas": len(rows),
        "errores_registrados": len(errors),
        "nota": (
            "La fecha usada es la fecha de emisión impresa en el informe. "
            "Los PDF sin texto legible pueden requerir OCR y se registran en errores_extraccion.csv."
        ),
    }
    write_outputs(args.salida, rows, errors, summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"Excel: {args.salida / 'tabla3_arbovirus_PR.xlsx'}")
    print(f"CSV: {args.salida / 'tabla3_arbovirus_PR.csv'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())